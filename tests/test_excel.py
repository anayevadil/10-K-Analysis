import io

import openpyxl
import pytest

from tenk.export.excel import workbook_bytes, write_workbook


@pytest.fixture
def workbook(statements):
    return openpyxl.load_workbook(io.BytesIO(workbook_bytes(statements)))


def rows_by_label(sheet):
    return {row[0].value: row for row in sheet.iter_rows(min_row=2) if row[0].value}


def test_sheets(workbook):
    assert workbook.sheetnames == [
        "Summary",
        "Income Statement",
        "Balance Sheet",
        "Cash Flow",
        "Ratios",
        "Sources",
    ]


def test_summary(workbook):
    sheet = workbook["Summary"]
    assert sheet["A1"].value == "Example Corp"
    assert sheet["B4"].value == "2020-12-31 to 2024-12-31"


def test_headers(workbook):
    header = [c.value for c in workbook["Income Statement"][1]]
    assert header[1] == "FY ending 2020-12-31"
    assert header[-1] == "FY ending 2024-12-31"


def test_values_in_millions(workbook):
    rows = rows_by_label(workbook["Income Statement"])
    assert rows["Revenue"][5].value == 1300
    assert rows["Diluted EPS"][5].value == pytest.approx(1.48)
    assert rows["Diluted shares"][5].value == 100


def test_subtotals_are_bold(workbook):
    rows = rows_by_label(workbook["Income Statement"])
    assert rows["Net income"][0].font.bold
    assert not rows["Research and development"][0].font.bold


def test_ratios_are_live_formulas(workbook):
    rows = rows_by_label(workbook["Ratios"])
    income = rows_by_label(workbook["Income Statement"])
    revenue_row = income["Revenue"][0].row
    gross_row = income["Gross profit"][0].row
    assert rows["Gross margin"][5].value == (
        f"=IFERROR('Income Statement'!F{gross_row}/'Income Statement'!F{revenue_row},\"\")"
    )


def test_first_year_growth_is_blank(workbook):
    rows = rows_by_label(workbook["Ratios"])
    assert rows["Revenue growth"][1].value is None
    assert rows["Revenue growth"][2].value.startswith("=IFERROR(")


def test_roe_uses_average_equity_after_first_year(workbook):
    rows = rows_by_label(workbook["Ratios"])
    assert "AVERAGE(" not in rows["Return on equity"][1].value
    assert "AVERAGE(" in rows["Return on equity"][5].value


def test_ratio_highlighting(workbook):
    ranges = list(workbook["Ratios"].conditional_formatting)
    assert len(ranges) == 8  # one per ratio, from the second year on
    assert str(ranges[0].sqref) == "C2:F2"
    assert all(len(r.rules) == 2 for r in ranges)  # better and worse


def test_balance_sheet_check_row(workbook):
    rows = rows_by_label(workbook["Balance Sheet"])
    check = rows["Check: assets - liabilities - equity"]
    assert check[1].value.startswith("='Balance Sheet'!B")


def test_sources(workbook):
    rows = rows_by_label(workbook["Sources"])
    assert rows["Revenue"][1].value == "SalesRevenueNet"
    assert rows["Gross profit"][1].value == "derived"
    assert rows["Gross profit"][5].value == "GrossProfit"


def test_write_to_path(statements, tmp_path):
    path = tmp_path / "model.xlsx"
    write_workbook(statements, path)
    assert openpyxl.load_workbook(path).sheetnames[0] == "Summary"
